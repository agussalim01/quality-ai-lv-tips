import streamlit as st
import pandas as pd
import plotly.express as px
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import io
import re
from datetime import datetime

st.set_page_config(page_title="Quality AI PRO - Landscape Bottom Approval", page_icon="📄", layout="wide")

# ================= TAB 1: MSA =================
def msa_tab():
    st.subheader("📊 MSA Gage R&R Dashboard")
    uploaded = st.file_uploader("Upload Data MSA (.csv)", type=["csv"])
    if uploaded:
        df = pd.read_csv(uploaded)
    else:
        df = pd.DataFrame({
            "Part": [1,1,1,2,2,2,3,3,3]*3,
            "Operator": ["A"]*9 + ["B"]*9 + ["C"]*9,
            "Measurement": [10.1,10.12,10.11,10.2,10.22,10.21,10.0,10.02,10.01,
                            10.11,10.13,10.12,10.21,10.23,10.22,10.01,10.03,10.02,
                            10.09,10.11,10.10,10.19,10.21,10.20,9.99,10.01,10.0]
        })
        st.caption("Data Demo - Upload CSV untuk data asli")
    c1, c2 = st.columns(2)
    with c1:
        st.dataframe(df.head(15), use_container_width=True)
    with c2:
        fig = px.histogram(df, x="Measurement", color="Operator", title="MSA Distribution")
        st.plotly_chart(fig, use_container_width=True)
    st.success("MSA Dashboard Aktif")

# ================= TAB 2: WI LANDSCAPE - APPROVAL SELALU DI BAWAH =================
def create_wi_landscape_bottom(template_file, doc_no, product, temp_no, rev, proc, part_req, tools_req, time_df, steps_text):
    # Load template asli
    doc = Document(template_file)
    
    # === 1. SIMPAN 2 TABEL ASLI, LALU HAPUS DARI ATAS ===
    # Simpan referensi tabelnya dulu
    saved_tables = list(doc.tables)  # ada 2 tabel: Approval & History
    
    # Hapus dari atas (biar WI bisa di atas)
    for tbl in saved_tables:
        tbl._tbl.getparent().remove(tbl._tbl)
    
    # Bersihkan paragraf kosong sisa template
    for p in list(doc.paragraphs):
        if not p.text.strip():
            p._p.getparent().remove(p._p)

    # === 2. TAMBAH WI DI ATAS DULU ===
    # Judul besar tengah
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"WORK INSTRUCTION\n{proc}\nProduct: {product}\nDocument No: {doc_no}")
    run.bold = True
    run.font.size = Pt(14)

    # Info Box
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Document No: {doc_no} | Temp: {temp_no} | Rev: {rev} | Release: {datetime.now().strftime('%d/%m/%Y')} | Process: {proc}")
    run.font.size = Pt(8)
    run.bold = True

    # Tabel Langkah Kerja
    p = doc.add_paragraph()
    run = p.add_run("Langkah Kerja")
    run.bold = True
    run.font.size = Pt(11)

    t2 = doc.add_table(rows=1, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        t2.style = 'Table Grid'
    except:
        t2.style = 'Grid Table Light'

    t2.cell(0,0).text = "Langkah Kerja"
    t2.cell(0,1).text = "Time (min)"
    for cell in t2.rows[0].cells:
        for par in cell.paragraphs:
            for r in par.runs: r.bold = True

    total = 0
    for _, row in time_df.iterrows():
        rw = t2.add_row().cells
        rw[0].text = str(row['step'])
        rw[1].text = str(row['time'])
        try: total += int(row['time'])
        except: pass
    rw = t2.add_row().cells
    rw[0].text = "Total"
    rw[1].text = str(total)
    for c in rw:
        for par in c.paragraphs:
            for r in par.runs: r.bold = True

    doc.add_paragraph("")

    # Part & Tools
    p = doc.add_paragraph()
    r = p.add_run("Part required :")
    r.bold = True
    for it in part_req.split('\n'):
        if it.strip(): doc.add_paragraph(f"●  {it.strip()}")

    p = doc.add_paragraph()
    r = p.add_run("Tools required :")
    r.bold = True
    for it in tools_req.split('\n'):
        if it.strip(): doc.add_paragraph(f"●  {it.strip()}")

    doc.add_paragraph("")

    # Detail Step
    pattern = r"(Step\s*\d+\s*:.*)"
    parts = re.split(pattern, steps_text, flags=re.I)
    steps = []
    for i in range(1, len(parts), 2):
        title = parts[i].strip()
        body = parts[i+1] if i+1 < len(parts) else ""
        bullets = [x.strip() for x in body.split('\n') if x.strip()]
        steps.append((title, bullets))
    if not steps:
        steps = [("Work Steps", [x for x in steps_text.split('\n') if x.strip()])]

    for title, bullets in steps:
        p = doc.add_paragraph()
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(10)
        for b in bullets:
            if b:
                clean = re.sub(r'^[•\-\*]\s*', '', b)
                doc.add_paragraph(f"●  {clean}")

    # === 3. PINDAHKAN 2 TABEL ASLI KE BAWAH (SELALU DI BAWAH) ===
    doc.add_page_break()
    
    for tbl in saved_tables:
        # Tempel balik tabel asli plek 100% ke bawah
        doc._body._element.append(tbl._tbl)
        doc.add_paragraph("")  # spasi

    # Update isi History (tabel terakhir) setelah dipindah ke bawah
    try:
        hist_table = doc.tables[-1]  # tabel history yang sudah di bawah
        if len(hist_table.rows) > 2:
            hist_table.cell(2,1).text = datetime.now().strftime("%m/%d/%Y")
            hist_table.cell(2,2).text = rev
            hist_table.cell(2,3).text = f"{proc} - {product}"
    except:
        pass

    return doc

def wi_tab():
    st.subheader("📄 WI Generator - Approval Selalu di Bawah")
    st.success("✅ Template landscape TMP-Common-General-0002 - APPROVAL & HISTORY otomatis di BAWAH")

    template_file = st.file_uploader("1. Upload Template Landscape (.docx) - TMP-Common-General-0002", type=["docx"], key="templ")
    
    c1, c2 = st.columns(2)
    with c1:
        doc_no = st.text_input("Document No", "WIN-13-005-0015", key="doc2")
        product = st.text_input("Product", "NT8", key="prod2")
        temp_no = st.text_input("Temp Number", "TMP-11-001-0002 rev.1", key="temp2")
        rev = st.text_input("Rev", "0", key="rev2")
        proc = st.text_input("Process Desc", "QC 100% Visual Inspection LV Tips", key="proc2")
    with c2:
        part_req = st.text_area("Part Required", "1 batch LV Tips after cutting", height=80, key="part2")
        tools_req = st.text_area("Tools Required", "Lampu LED 6 watt 580 LM\nTatakan LV tips\nPinset\nIonizer", height=120, key="tool2")

    time_df = st.data_editor(pd.DataFrame([
        {"step": "Step 1 :  Susun Tips", "time": 10},
        {"step": "Step 2 :  Check visual defect", "time": 80},
        {"step": "Step 3 :  Data Tips OK dan Tips Reject", "time": 5},
        {"step": "Step 4 :  Packing LV Tips", "time": 5},
        {"step": "Step 5 : Manual Filling Process", "time": 30},
        {"step": "Step 6 : 2nd Visual Check", "time": 5},
        {"step": "Step 7 : Ionizer", "time": 5},
    ]), num_rows="dynamic", use_container_width=True, key="time_land")

    raw = st.text_area("Detail Langkah (format Step X :)", 
"""Step 1 :  Susun Tips
Susun LV tips di atas wadah (tatakan) LV tips secara horizontal dengan menggunakan pinset.
Step 2 :  Check visual defect
Cek tips secara kasat mata dengan pencahayaan lampu LED dengan jarak 30cm
Cek visual bagian ujung sampai pangkal tips
Pastikan tidak terdapat bubble, sinkmark, penyok, short, flashing, kontaminasi
Step 3 :  Data Tips OK dan Tips Reject
Catat jumlah tips OK dan tips yang Reject hasil QC 100% di form QC 100%
Step 4 :  Packing LV Tips
Masukkan kembali tips yang OK ke dalam plastik ziplock sebelumnya.""", height=220, key="raw_land")

    if st.button("✨ Generate WI - Approval di Bawah", type="primary", use_container_width=True):
        if not template_file:
            st.error("Upload dulu template landscape TMP-Common-General-0002!")
            return
        doc = create_wi_landscape_bottom(template_file, doc_no, product, temp_no, rev, proc, part_req, tools_req, time_df, raw)
        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)
        st.success(f"WI {doc_no} jadi! WI di atas, APPROVAL & HISTORY di bawah - Plek template asli!")
        st.download_button("📥 Download WI Landscape (Approval di Bawah)", data=bio, file_name=f"{doc_no}_Landscape_ApprovalBawah_{product}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)

# MAIN
st.title("📊 Quality AI PRO - MSA + WIN Landscape Bottom")
tab1, tab2 = st.tabs(["📊 MSA Dashboard", "📄 WI Generator Landscape"])

with tab1:
    msa_tab()
with tab2:
    wi_tab()
